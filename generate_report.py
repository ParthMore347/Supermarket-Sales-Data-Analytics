"""Generate charts and a Word executive report for the Supermarket Sales project."""

from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


DATA_URL = (
    "https://raw.githubusercontent.com/aungpyaeap/supermarket-sales/master/"
    "supermarket_sales%20-%20Sheet1.csv"
)
WORKSPACE_ROOT = Path(__file__).resolve().parent
IMAGE_DIRECTORY = WORKSPACE_ROOT / "report_images"
KPI_IMAGE_PATH = IMAGE_DIRECTORY / "kpi_summary.png"
PRODUCT_IMAGE_PATH = IMAGE_DIRECTORY / "revenue_by_product.png"
REPORT_PATH = WORKSPACE_ROOT / "ParthMore_ProjectReport.docx"


def generate_fallback_sales_data(row_count: int = 1000) -> pd.DataFrame:
    """Create an offline dataset with the fields required by this report."""
    rng = np.random.default_rng(42)
    product_lines = np.array(
        [
            "Health and beauty",
            "Electronic accessories",
            "Home and lifestyle",
            "Sports and travel",
            "Food and beverages",
            "Fashion accessories",
        ]
    )
    dates = pd.date_range("2019-01-01", periods=89, freq="D")
    unit_prices = rng.uniform(10, 100, row_count).round(2)
    quantities = rng.integers(1, 11, row_count)
    totals = (unit_prices * quantities * rng.uniform(1.01, 1.10, row_count)).round(2)
    return pd.DataFrame(
        {
            "Invoice ID": [f"SYN-{index:04d}" for index in range(row_count)],
            "Product line": rng.choice(product_lines, row_count),
            "Total": totals,
            "Date": rng.choice(dates, row_count),
            "Rating": rng.uniform(5, 10, row_count).round(1),
        }
    )


def load_sales_data() -> pd.DataFrame:
    """Download the published CSV, with a visible offline fallback."""
    request = Request(DATA_URL, headers={"User-Agent": "supermarket-sales-report/1.0"})
    try:
        with urlopen(request, timeout=30) as response:
            return pd.read_csv(BytesIO(response.read()))
    except (HTTPError, URLError, TimeoutError) as error:
        print(
            f"Unable to download the requested dataset ({error}). "
            "Using a reproducible synthetic fallback dataset."
        )
        return generate_fallback_sales_data()


def clean_sales_data(sales_data: pd.DataFrame) -> pd.DataFrame:
    """Validate, deduplicate, and normalize fields needed by the report."""
    required_columns = {
        "Invoice ID",
        "Product line",
        "Total",
        "Date",
        "Rating",
    }
    missing_columns = required_columns.difference(sales_data.columns)
    if missing_columns:
        raise ValueError(f"Dataset is missing required columns: {sorted(missing_columns)}")

    cleaned_data = sales_data.drop_duplicates().copy()
    cleaned_data["Date"] = pd.to_datetime(cleaned_data["Date"], errors="coerce")
    cleaned_data["Total"] = pd.to_numeric(cleaned_data["Total"], errors="coerce")
    cleaned_data["Rating"] = pd.to_numeric(cleaned_data["Rating"], errors="coerce")
    cleaned_data = cleaned_data.dropna(
        subset=["Invoice ID", "Product line", "Total", "Date", "Rating"]
    )
    if cleaned_data.empty:
        raise ValueError("No valid rows remain after cleaning the dataset.")
    return cleaned_data


def compute_kpis(sales_data: pd.DataFrame) -> dict[str, float]:
    """Compute the four executive KPIs."""
    total_orders = sales_data["Invoice ID"].nunique()
    total_revenue = float(sales_data["Total"].sum())
    return {
        "Total Revenue": total_revenue,
        "Total Orders": float(total_orders),
        "Average Order Value": total_revenue / total_orders if total_orders else 0.0,
        "Average Rating": float(sales_data["Rating"].mean()),
    }


def save_kpi_chart(kpis: dict[str, float]) -> None:
    """Save a clean executive banner with four KPI cards."""
    figure, axes = plt.subplots(1, 4, figsize=(16, 3.6))
    figure.patch.set_facecolor("#F4F7FB")
    card_colors = ["#16324F", "#2A9D8F", "#E9C46A", "#E76F51"]
    display_values = [
        f"${kpis['Total Revenue']:,.0f}",
        f"{kpis['Total Orders']:,.0f}",
        f"${kpis['Average Order Value']:,.2f}",
        f"{kpis['Average Rating']:.2f}/10",
    ]
    for axis, label, value, color in zip(
        axes, kpis, display_values, card_colors, strict=True
    ):
        axis.set_facecolor(color)
        axis.text(
            0.5,
            0.62,
            value,
            ha="center",
            va="center",
            color="white",
            fontsize=18,
            fontweight="bold",
        )
        axis.text(
            0.5,
            0.25,
            label,
            ha="center",
            va="center",
            color="white",
            fontsize=10,
        )
        axis.set_xticks([])
        axis.set_yticks([])
        for spine in axis.spines.values():
            spine.set_visible(False)
    figure.suptitle(
        "Supermarket Sales — Executive KPI Summary",
        fontsize=16,
        fontweight="bold",
        color="#16324F",
    )
    figure.tight_layout()
    figure.savefig(KPI_IMAGE_PATH, dpi=180, bbox_inches="tight")
    plt.close(figure)


def save_product_revenue_chart(sales_data: pd.DataFrame) -> pd.DataFrame:
    """Save a ranked Seaborn product-line revenue chart and return its summary."""
    product_revenue = (
        sales_data.groupby("Product line", as_index=False)["Total"]
        .sum()
        .sort_values("Total", ascending=False)
    )
    figure, axis = plt.subplots(figsize=(10, 6))
    sns.barplot(
        data=product_revenue,
        y="Product line",
        x="Total",
        hue="Product line",
        palette="crest",
        legend=False,
        ax=axis,
    )
    axis.set_title("Revenue by Product Line", fontsize=16, fontweight="bold")
    axis.set_xlabel("Total Revenue ($)")
    axis.set_ylabel("")
    axis.xaxis.set_major_formatter(
        plt.FuncFormatter(lambda value, _: f"${value:,.0f}")
    )
    sns.despine()
    figure.tight_layout()
    figure.savefig(PRODUCT_IMAGE_PATH, dpi=180, bbox_inches="tight")
    plt.close(figure)
    return product_revenue


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.style.font.name = "Aptos Display"


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def create_word_report(
    sales_data: pd.DataFrame, kpis: dict[str, float], product_revenue: pd.DataFrame
) -> None:
    """Build the executive Word report with charts and decision statements."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_heading("Supermarket Sales Executive Decision Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "Author: Parth More | IBM SkillsBuild Data Analytics with AI Internship"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.italic = True

    add_heading(document, "1. Executive KPIs (Level 1)")
    kpi_table = document.add_table(rows=1, cols=2)
    kpi_table.style = "Light Shading Accent 1"
    kpi_table.rows[0].cells[0].text = "Metric"
    kpi_table.rows[0].cells[1].text = "Value"
    formatted_kpis = [
        ("Total Revenue", f"${kpis['Total Revenue']:,.2f}"),
        ("Total Orders", f"{kpis['Total Orders']:,.0f}"),
        ("Average Order Value (AOV)", f"${kpis['Average Order Value']:,.2f}"),
        ("Average Rating", f"{kpis['Average Rating']:.2f} / 10"),
    ]
    for metric, value in formatted_kpis:
        cells = kpi_table.add_row().cells
        cells[0].text = metric
        cells[1].text = value
    document.add_picture(str(KPI_IMAGE_PATH), width=Inches(7.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "2. Revenue Drivers (Level 3)")
    top_product = product_revenue.iloc[0]
    bottom_product = product_revenue.iloc[-1]
    document.add_paragraph(
        f"{top_product['Product line']} is the leading product line, generating "
        f"${top_product['Total']:,.2f}. {bottom_product['Product line']} is the "
        f"lowest-revenue category at ${bottom_product['Total']:,.2f}."
    )
    document.add_picture(str(PRODUCT_IMAGE_PATH), width=Inches(6.8))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "3. Risk Analysis (Level 4)")
    category_ratings = (
        sales_data.groupby("Product line", as_index=False)["Rating"]
        .mean()
        .sort_values("Rating")
    )
    overall_rating = sales_data["Rating"].mean()
    low_rated_categories = category_ratings[
        category_ratings["Rating"] < overall_rating
    ]
    if low_rated_categories.empty:
        document.add_paragraph(
            f"No product category is below the overall rating benchmark of "
            f"{overall_rating:.2f}/10. Continue monitoring category ratings."
        )
    else:
        document.add_paragraph(
            f"The overall customer rating is {overall_rating:.2f}/10. "
            "Categories below this benchmark represent potential customer "
            "experience and churn risks:"
        )
        for _, row in low_rated_categories.iterrows():
            add_bullet(
                document,
                f"{row['Product line']}: {row['Rating']:.2f}/10 average rating.",
            )

    add_heading(document, "4. Strategic Recommendations")
    recommendation_text = [
        (
            f"Fact: {top_product['Product line']} leads revenue at "
            f"${top_product['Total']:,.2f}. "
            "Insight: demand is concentrated in a category leader. "
            "Opportunity: increase basket size through complementary products. "
            "Action: test cross-sell bundles and track AOV for four weeks."
        ),
        (
            f"Fact: {bottom_product['Product line']} is the lowest-revenue "
            f"category at ${bottom_product['Total']:,.2f}. "
            "Insight: visibility, assortment, or pricing may be limiting demand. "
            "Opportunity: recover revenue with a targeted intervention. "
            "Action: run an end-cap or price test against a control period."
        ),
        (
            f"Fact: {len(low_rated_categories)} product categories are below "
            f"the overall rating benchmark of {overall_rating:.2f}/10. "
            "Insight: category-level service or product issues may affect repeat purchase. "
            "Opportunity: address root causes before churn grows. "
            "Action: collect category-specific feedback and set a rating recovery target."
        ),
    ]
    for statement in recommendation_text:
        add_bullet(document, statement)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Generated by generate_report.py")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)
    document.save(REPORT_PATH)


def main() -> None:
    IMAGE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    raw_sales_data = load_sales_data()
    sales_data = clean_sales_data(raw_sales_data)
    kpis = compute_kpis(sales_data)
    save_kpi_chart(kpis)
    product_revenue = save_product_revenue_chart(sales_data)
    create_word_report(sales_data, kpis, product_revenue)
    print(f"Report generated successfully: {REPORT_PATH}")
    print(f"Charts saved successfully in: {IMAGE_DIRECTORY}")


if __name__ == "__main__":
    main()
